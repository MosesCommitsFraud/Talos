from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path('C:/DEV/Talos')
OUT=ROOT/'output/Talos_Hardware_Modelle_September_2026_Cowork.docx'
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.65)
s.left_margin=Inches(.72); s.right_margin=Inches(.72)
s.footer_distance=Inches(.28)
for name,size in [('Normal',10.5),('Title',27),('Subtitle',13),('Heading 1',17),('Heading 2',12)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.05
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(11)
for el in doc.styles.element.xpath('.//w:pBdr'):el.getparent().remove(el)
doc.styles['Normal'].paragraph_format.widow_control=True
p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Talos  |  September 2026  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
def p(t='',style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(t):
 doc.add_page_break();doc.add_heading(t,1)
def table(headers,rows,widths=None,size=9.5):
 t=doc.add_table(rows=0,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=widths or [6.83/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate([headers]+rows):
  cells=t.add_row().cells
  trpr=t.rows[-1]._tr.get_or_add_trPr(); avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,(c,txt,w) in enumerate(zip(cells,row,widths)):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   par=c.paragraphs[0];par.paragraph_format.space_before=Pt(1.4 if headers[0]=='Begriff' else 3);par.paragraph_format.space_after=Pt(1.4 if headers[0]=='Begriff' else 3);par.paragraph_format.line_spacing=1.0
   r=par.add_run(str(txt));r.font.size=Pt(size)
   if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'24445A' if i==0 else ('F0F4F7' if i%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
 gap=p('');gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.line_spacing=Pt(3);gap.add_run().font.size=Pt(3)
 return t
def link(par,label,url):
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245B80');pr.append(c);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);el.append(r);par._p.append(el)

p('Talos Hardware und Modelle','Title')
p('Plattformvergleich Software Leistung und Erweiterungen','Subtitle')
p('Stand 7. September 2026')
p('Für einen lokalen Talos-Betrieb sind drei Entscheidungen maßgeblich: Welche Aufgaben soll das Modell zuverlässig lösen, wie viele Anfragen sollen gleichzeitig laufen und welche Antwortzeit ist akzeptabel? Daraus ergeben sich Modell, Speicher und Rechenplattform. Die größte Parameterzahl oder der höchste beworbene Rechenwert allein liefert keine verlässliche Auswahl.')
h('Die wichtigsten Entscheidungen')
table(['Ziel','Sinnvoller Ausgangspunkt'],[
('Kompakter Einstieg mit vorhandener NVIDIA-Umgebung','Ein DGX Spark oder vergleichbares GB10-System mit Qwen3.6-35B-A3B beziehungsweise Qwen3.8-27B.'),
('Größere Modelle bei überschaubarem Ausbau','Zwei Sparks für DeepSeek-V4-Flash oder GLM-5.3-Flash mit passender Quantisierung. Vier Geräte erst durch Kapazitäts- und Lasttests begründen.'),
('Schnelle Antworten und mehr gleichzeitige Arbeit','Custom-Server mit RTX PRO 6000 Blackwell prüfen; 96 GB schneller GPU-Speicher pro Karte, aber deutlich höherer Infrastrukturbedarf.'),
('Viel Speicher in einem einzelnen Desktop','Apple Ultra mit 256 bis 512 GB prüfen; AMD mit 128 beziehungsweise bis zu 192 GB als Alternative. Laufzeit und Modellunterstützung separat bewerten.'),
('Sehr große Modelle wie GLM-5.3','Große Speicherplattform oder mehrere Beschleuniger. Vier Sparks sind ein bedingter Quantisierungsfall, kein pauschal freigegebenes Produktivsystem.')],[2.1,4.73])
h('So sind die Angaben zu lesen')
p('Spezifikation bedeutet Herstellerangabe. Messbericht bedeutet veröffentlichter Lauf auf konkret benannter Hardware, nicht eine eigene Talos-Messung. Planung bedeutet rechnerische Abschätzung oder vorgeschlagene Konfiguration. Diese Kategorien werden im Dokument getrennt. Geschwindigkeitstabellen bilden keine einheitliche Rangliste, wenn Quantisierung, Kontext oder Testverfahren abweichen.')
p('Die Modellqualität entsteht aus Modellversion, Quantisierung und vollständigem Arbeitsablauf. Dieselben Gewichte werden auf einer teureren GPU nicht automatisch klüger. Die GPU kann jedoch kürzere Wartezeiten, höhere Präzision oder größere Kontexte ermöglichen.')
h('Orientierung')
p('Nachrichtenstand und Hardware → lokale Leistung und Artificial Analysis → Software in Talos → Hardware Upgrades → Glossar und Quellen. Quellenkennungen wie [S1] verweisen auf die verlinkten Originalquellen am Ende.')

page('1 Nachrichtenstand im September 2026')
p('Stichtag ist der 7. September 2026. Ankündigungen für spätere Termine sind als solche ausgewiesen. Bei Modellen sind API-Verfügbarkeit, offene Gewichte und die Unterstützung durch eine lokale Laufzeit drei getrennte Meilensteine.')
table(['Datum oder Stand','Neue Entwicklung','Bedeutung für Talos'],[
('IFA Anfang September','NVIDIA kündigt RTX Spark für Oktober an. Bis 128 GB Unified Memory; Windows-Plattform. [S29]','Zusätzliche Desktop- und Mobiloption. Keine Gleichsetzung mit dem DGX-Spark-Linux-Stack und keine Übernahme seiner Benchmarks.'),
('25. August 2026','Apple stellt M5 Max und M5 Ultra im Mac Studio vor. M5 Ultra bis 512 GB und 1,2 TB/s; angekündigte Verfügbarkeit ab 22. September. [S6]','Großer Einzelrechner als Alternative zum Cluster. Am Stichtag keine hier qualifizierten Messreihen für die Zielmodelle auf Seriengeräten.'),
('26. August 2026','Qwen3.8-Flash-Next mit neuer Architektur und zusätzlichen N-Gram-Tabellen. [S14]','Mehr Modellkapazität bei wenig aktiven Parametern; Modellunterstützung und Ablage der Tabellen werden entscheidend.'),
('Ende August 2026','GLM-5.3 und GLM-5.3-Flash liegen als offene Gewichte vor. Frühe GB10-Deployments und neue NIM-Anleitung verfügbar. [S15–S16, S31]','Große Modelle werden lokal interessanter. Funktionsumfang jeder Quantisierung prüfen, insbesondere Bildverarbeitung und Werkzeuge.'),
('31. Juli bis 21. August','DeepSeek aktualisiert Flash und Pro; Flash-Vision-Exp startet am 21. August als experimentelles API-Modell. [S28]','API-Updates verbessern nicht automatisch lokal gespeicherte Gewichte. Vision-Exp ist nicht das lokale Flash-Textmodell.'),
('Q3 laut AMD-Ankündigung','Ryzen AI Max PRO 400 mit bis 192 GB, OEM-Systeme von HP und Lenovo angekündigt. [S8]','Mehr Kapazität als 128-GB-Geräte. Lieferbarkeit und GPU-Speicherbudget der konkreten Konfiguration bestätigen lassen.'),
('Juni und Juli 2026','DGX Sync Cluster Assistant und verbessertes Speicherfehlermanagement. [S2]','Clusteraufbau wird einfacher; verteilte Modellinferenz braucht weiterhin passende Laufzeit und Netzwerk.'),
('Juli bis Anfang September','Huawei zeigt Atlas 950 SuperPoD; Moore Threads meldet GLM-5.3-Flash-Anpassung am 27. August und weitere Inferenzkooperation am 4. September. [S10, S11a]','Relevante Entwicklung im chinesischen Serverökosystem; keine belegte Austauschbarkeit mit einem Spark.')],[1.12,2.9,2.81],9.2)
p('Die Veröffentlichungen verändern vor allem die Auswahl an Laufzeiten und Modellformaten. Neue Spitzenwerte gelten nur mit der dazugehörigen Konfiguration; für eine Beschaffung bleibt ein repräsentativer Talos-Test notwendig.')

page('2 Kompakte Hardware im Vergleich')
table(['Plattform','Speicher und Bandbreite','Einordnung'],[
('NVIDIA DGX Spark','128 GB gemeinsam; 273 GB/s; GB10 mit Arm-CPU und Blackwell-GPU. [S1]','Guter Anschluss an CUDA und vorhandene Talos-Rezepte. 1 oder 4 TB SSD je Variante. RAM nicht wie DIMM-Speicher erweiterbar.'),
('ASUS Ascent GX10 und andere GB10-OEMs','GX10 ebenfalls 128 GB und 273 GB/s. [S3]','Gleiche Chipklasse, aber Kühlung, SSD, Garantie und Updatezeitpunkt vergleichen. OEM-Wechsel allein ist kein Rechensprung.'),
('AMD Ryzen AI Max+ 395','Bis 128 GB gemeinsamer Speicher; typische LPDDR5x-8000-Konfiguration rechnerisch 256 GB/s. [S7]','Kompakter x86-Rechner. ROCm beziehungsweise Vulkan statt CUDA; tatsächliche GPU-Zuweisung und Betriebssystem prüfen.'),
('AMD Ryzen AI Max+ PRO 495','Bis 192 GB; 256-Bit LPDDR5x-8533 ergibt rechnerisch rund 273 GB/s. [S8–S9]','Kapazitätsupgrade. AMD nennt bis 160 GB dedizierbaren Grafikspeicher; 192 GB sind nicht vollständig als GPU-Budget zu verplanen.'),
('Apple Mac Studio M3 Ultra','Referenzgeneration bis 512 GB; 819 GB/s. [S5]','Gemessene lokale MLX-Läufe verfügbar. RAM bei Kauf festlegen; bestehende Geräte und verfügbare Händlerkonfiguration unterscheiden.'),
('Apple Mac Studio M5 Ultra','Angekündigt bis 512 GB und 1,2 TB/s. [S6]','Auslieferung laut Apple ab 22. September. Interessanter Ausbaupfad; keine lineare Hochrechnung von M3-Messungen.'),
('Apple Max statt Ultra','M5 Max angekündigt bis 128 GB und 614 GB/s. [S6]','Kleinere Modelle und kompakter Betrieb; kein Ersatz für die Speicherkapazität eines Ultra mit 256/512 GB.')],[1.68,2.42,2.73],9.4)
h('Was der Spark praktisch bietet')
p('Das Gerät hat 10-GbE und ConnectX-7-Anschlüsse. NVIDIA nennt 140 W für den GB10-Chip und ein 240-W-Netzteil. Das sind keine gemessenen Stromkosten. Der beworbene Petaflop gilt für FP4 mit Sparsity; er lässt sich weder mit beliebigen TOPS-Angaben noch direkt mit Tokens pro Sekunde vergleichen. [S1]')
h('Warum die Plattformentscheidung mehr als Speicher ist')
p('NVIDIA, AMD und Apple verwenden unterschiedliche Beschleunigungssoftware. Talos kann einen Modellserver über das Netzwerk ansprechen; dadurch muss die Anwendung nicht auf derselben Hardware wie das Modell laufen. Auf einem Mac sollte die Metal-/MLX-Inferenz nativ erfolgen, während die Talos-Dienste getrennt oder in einer geeigneten Containerumgebung betrieben werden. Ein CUDA-Container lässt sich nicht unverändert auf Apple- oder AMD-Grafik übertragen.')

page('3 Custom Server und chinesische Plattformen')
table(['Beschleuniger','Speicher je Einheit','Praktische Konsequenz'],[
('RTX 6000 Ada Generation','48 GB GDDR6; 960 GB/s [S4]','Ältere Generation. Zwei Karten ergeben nominell 96 GB verteilt; kein NVFP4-Blackwell-Ersatz.'),
('RTX PRO 6000 Blackwell','96 GB GDDR7 ECC; 1.792 GB/s [S12]','Workstation bis 600 W, Max-Q 300 W, Server 400–600 W je Karte. Edition und Kühlung verbindlich spezifizieren.'),
('AMD Radeon AI PRO R9700','32 GB GDDR6 [S30]','Einstieg über diskrete GPU. Mehrere Karten verlangen passende Plattform und geprüfte Peer-to-Peer-Kommunikation.'),
('AMD Instinct MI350X','288 GB HBM3E; 8 TB/s [S13]','Rechenzentrumsklasse. Modularer Serverbeschleuniger, keine gewöhnliche Desktop-Steckkarte.'),
('Moore Threads MTT S5000','80 GB; 1,6 TB/s laut Herstellerdokumentation [S11]','MUSA-Softwareumgebung und MTLink. Hersteller nennt vLLM-/SGLang-Unterstützung; konkretes Modell und Quantisierung abnehmen.'),
('Huawei Ascend und Atlas','Konfiguration abhängig von Karte und Serversystem [S10]','CANN-Ökosystem; GLM nennt auch Ascend-Laufzeiten. Atlas 950 SuperPoD ist eine große Clusterplattform, kein Mini-PC. [S16]')],[1.87,2.05,2.91],9.4)
h('Drei beispielhafte Serverkonfigurationen')
p('Planung A: Eine RTX PRO 6000 Blackwell mit 96 GB, 16–32 CPU-Kernen, 128–256 GB System-RAM, 2–4 TB NVMe und 10-GbE. Geeignet als schneller Modellserver für die Qwen-Modelle der 27/35B-Klasse. Bei BF16 bleiben weniger Reserven für Hilfsmodelle und lange Kontexte; eine eigene CPU-Instanz für Talos entlastet das System.')
p('Planung B: Zwei RTX PRO 6000 mit 192 GB verteiltem VRAM, Threadripper-PRO- oder EPYC-Plattform, 256–512 GB RAM und 4–8 TB NVMe. DeepSeek-V4-Flash mit passendem kompaktem Checkpoint ist ein Kandidat; GLM-Flash mit etwa 182 GB Gewichten wäre sehr knapp. Für Letzteres besser vier Karten oder andere Speicheraufteilung prüfen.')
p('Planung C: Vier RTX PRO 6000 mit 384 GB VRAM, 512 GB bis 1 TB RAM, 8 TB NVMe und zertifiziertem Mehr-GPU-Gehäuse. Für große Flash-Modelle mit mehr Cache und Parallelität. GLM-5.3 in 4 Bit liegt bereits bei etwa 375 GB reinen Gewichten: vier Karten reichen dafür nicht komfortabel. Acht Karten oder eine größere HBM-Plattform sind die robustere Planungsrichtung.')
p('Bei vier 600-W-Karten entstehen bis zu 2,4 kW reine GPU-Leistungsaufnahme. Netzteil, Stromkreis, Luftführung und Kühlung müssen zusätzlich CPU und Peripherie tragen. PCIe-Steckplätze müssen elektrisch ausreichend angebunden sein; sichtbare x16-Slots garantieren keine x16-Verbindung. Mehr-GPU-RTX ist kein NVLink-HBM-System. Alle Konfigurationen sind Vorschläge, keine gemessenen Stücklisten.')

page('4 Lokale Geschwindigkeit im Vergleich')
p('Veröffentlichte Fremdmessungen, keine eigenen Talos-Messungen. Die Werte beschreiben unterschiedliche Laufzeiten, Quantisierungen und Aufgaben; sie erlauben eine Größenordnung, aber keine kontrollierte Hardware-Rangliste.')
table(['Modell und Hardware','Gemessene Ausgabe','Testbedingungen und Quelle'],[
('Qwen3.6-35B · 1 Spark','67,4 t/s','Q4_K_M, llama-bench, ein Stream. [S20]'),
('Qwen3.8-27B · 1 GB10','24,7 t/s','vLLM mit MTP; Median aus sechs gemischten Aufgaben, davon drei deutsch. [S18]'),
('Qwen3.8-27B · AMD 395','14,0 / 30,6–36,0 t/s','Ohne / mit Greedy-MTP; ROCm FP4_FAST. Gewinn stark samplingabhängig. [S23]'),
('Qwen3.8-27B · M3 Ultra 256 GB','43,4 t/s','MLX 4 Bit, adaptives MTP, 8K Eingabe, Thinking aus; TTFT 24,7 s. [S24]'),
('Qwen3.8-27B · RTX PRO 6000','150 roh / 84 sichtbar t/s','Max-Q 300 W, NVFP4, vLLM mit MTP; Denk-Tokens erklären die Differenz. [S26]'),
('Qwen3.8-Flash-Next · 1 GB10','20,8–31,1 t/s','NVFP4, vLLM MTP 2; Prosa bis Code/Reasoning. [S18]'),
('Qwen3.8-Flash-Next · M3 Ultra','23,0 t/s','256 GB, MLX 4 Bit ohne MTP, 8K Eingabe; TTFT 9,4 s. [S24]'),
('DeepSeek-V4-Flash · 1 / 2 Sparks','ca. 14 / 38,5 t/s','1 Gerät: IQ2 + ds4; 2: gemischter offizieller Checkpoint, vLLM TP2 + MTP. [S20]'),
('GLM-5.3-Flash · 2 Sparks','19,6–30,3 t/s','NVFP4, vLLM MTP 5, Textmodus, Temperatur 0; Prosa bis strukturierte Ausgabe. [S21]'),
('GLM-5.3-Flash · 4 Sparks','31,6 t/s einzeln','NVFP4, vLLM TP4 + DFlash2, gemischte 4K-Prompts; TTFT 2,75 s. [S22]'),
('GLM-5.3-Flash · M3 Ultra','27,8 t/s','256 GB, MLX 4 Bit ohne MTP, 8K Eingabe; TTFT 22,8 s. [S24]')],[2.35,1.42,3.06],9.2)
h('Was sich daraus ableiten lässt')
p('Ein RTX-Server ist für schnelle Ausgabe ein plausibler Kandidat; ein Ultra bietet viel Speicher in einem Gerät. Zwei bis vier Sparks ermöglichen große Modelle mit dem vorhandenen NVIDIA-Stack. Vier Sparks liefern im GLM-Test bei vier parallelen Anfragen 64,1 t/s insgesamt, also etwa 16 t/s je Anfrage. Mehr Geräte bedeuten daher nicht proportional schnellere Einzelantworten. [S22]')
p('Für das volle GLM-5.3, M5 Ultra, PRO 495 und chinesische Plattformen liegt hier kein ausreichend dokumentierter, vergleichbarer lokaler Messwert vor. Die nächste Seite zeigt ergänzend unabhängige Modellbewertungen und API-Messungen von Artificial Analysis; deren Geschwindigkeiten sind keine Spark- oder RTX-Ergebnisse.')

page('5 Modellqualität mit Artificial Analysis')
p('Einheitliche Vergleichsbasis: Artificial Analysis Intelligence Index v4.2, abgerufen am 7. September 2026. Höher bedeutet besser im kombinierten Test; der Index ist keine Prozentquote richtiger Antworten. Ältere Indexversionen und Suchtreffer können andere Werte zeigen. [AA1–AA6]')
table(['Modell und Denkmodus','AA Index v4.2','AA API t/s','Lokale Einordnung'],[
('Qwen3.6-35B-A3B · Reasoning','26','124,9','Kompakter Einstieg auf einem Spark.'),
('Qwen3.8-27B · xhigh','41','47,2','Höherer Index bei kleiner Modellgröße; vorhandener Talos-Standard.'),
('Qwen3.8-Flash-Next','46','74,1','Mehr Qualität im Index; deutlich mehr Speicherbedarf.'),
('DeepSeek-V4-Flash 0731 · max','41','129,9','Hohe API-Ausgaberate; lokal auf zwei Sparks dokumentiert.'),
('GLM-5.3-Flash','46','56,1','Gleicher gerundeter Index wie Qwen Flash Next; auf zwei/vier Sparks gemessen.'),
('GLM-5.3 · max','49','77,7','Höchster Index dieser Auswahl; erheblich größere Hardware nötig.')],[2.13,.79,.86,3.05],9.5)
h('Dokumentenerstellung und professionelle Aufgaben')
p('GDPval-AA v2 bewertet reale Berufsaufgaben mit Werkzeugen und Webzugriff; die Ergebnisse werden blind paarweise verglichen. AA-Briefcase untersucht komplexe Wissensarbeit mit Tabellen, Präsentationen und Memos und berücksichtigt fachliche Erfüllung, Analyse und Darstellung. Beide berichten Elo-Werte, keine Prozentquoten. [AA7–AA8]')
table(['Modell','GDPval AA v2 Elo','AA Briefcase Elo'],[
('Qwen3.6-35B-A3B · Reasoning','993 ±14','—'),
('Qwen3.8-27B · xhigh','—','—'),
('Qwen3.8-Flash-Next','1.650 ±17','—'),
('DeepSeek-V4-Flash 0731 · max','1.470 ±15','1.259 −7/+7'),
('GLM-5.3-Flash','1.671 ±16','—'),
('GLM-5.3 · max','1.678 ±16','1.515 −8/+9')],[3.03,1.9,1.9],9.5)
p('Stand 7. September: ± beziehungsweise −/+ bezeichnet das veröffentlichte Unsicherheitsintervall. — = kein hier verifizierter Einzelwert für diesen Denkmodus. Der sichtbare Qwen-27B-Wert ohne Reasoning wird nicht xhigh zugeschrieben.')
p('Die GDPval-Intervalle von GLM-5.3 und GLM-Flash überlappen deutlich: kein klar belegter Vorteil des größeren Modells. Qwen Flash Next liegt nahe an beiden. Bei AA-Briefcase liegt GLM-5.3 deutlich vor DeepSeek-Flash.')

page('Wissen Recherche und Büroarbeit im Vergleich')
p('Für Talos zählen vollständige Arbeitsabläufe: Informationen finden, Quellen verstehen, Aussagen prüfen und daraus ein brauchbares Dokument erstellen. Die folgenden Werte stammen aus derselben Qwen-Vergleichstabelle; CoWorkBench ist ausdrücklich ein interner Anbietertest. Höher ist besser. [S14]')
table(['Benchmark und Schwerpunkt','Qwen 27B','Qwen Flash Next','DeepSeek Flash 0731'],[
('CoWorkBench · längere Büroaufgaben','70,7','73,9','45,1'),
('JobBench · berufliche Aufgaben','33,4','55,7','41,3'),
('Toolathlon Verified · Werkzeugabläufe','67,1','73,5','70,3'),
('IFBench · präzise Vorgaben einhalten','79,5','81,3','79,2'),
('HLE ohne Tools · breites schwieriges Wissen','30,8','35,9','33,8'),
('GPQA Diamond · wissenschaftliches Fachwissen','89,2','91,7','90,8')],[3.08,1.05,1.35,1.35],9.5)
p('Qwen Flash Next verbessert JobBench gegenüber Qwen 27B um 22,3 Punkte; bei CoWorkBench beträgt der Abstand nur 3,2 Punkte. DeepSeek liegt bei Toolathlon nahe an Qwen Flash Next, im internen CoWorkBench aber deutlich zurück. Das spricht für eine Auswahl nach Arbeitsaufgabe statt nach allgemeinem Modellrang.')
h('Welche Tests welchen Teil der Arbeit abdecken')
table(['Aufgabe','Passende Benchmarks und Bedeutung'],[
('Dokumente erstellen','AA-Briefcase und GDPval-AA: fertige Arbeitsergebnisse bewerten, einschließlich Qualität der Analyse und Präsentation. [AA7–AA8]'),
('Tabellen analysieren','AA-AnalystAgent: Zahlen aus Tabellen und Dokumenten herleiten. pass^5 verlangt fünf korrekte Wiederholungen, nicht nur einen erfolgreichen Versuch. [AA12]'),
('In Dokumenten recherchieren','AA-LCR v1.1: Informationen aus langen Texten verbinden. GDP.pdf: berufliche Dokumente verstehen und detaillierte Kriterien erfüllen. Das ist noch keine Bewertung des fertigen Word-Layouts. [AA9–AA10]'),
('Wissen zuverlässig nutzen','AA-Omniscience prüft Faktenwissen und Halluzinationen. HLE und GPQA ergänzen anspruchsvolles Wissen, messen aber keine aktuelle Webrecherche. [AA11, S14]'),
('Mehrstufige Arbeit erledigen','CoWorkBench, JobBench und Toolathlon bewerten längere Aufgaben beziehungsweise Werkzeuge; IFBench prüft genaue Ausgabevorgaben. [S14]')],[1.72,5.11],9.5)
p('Für AA-LCR, GDP.pdf, AA-Omniscience und AA-AnalystAgent wird hier kein nicht verifizierter Zielmodellwert ergänzt. Bei Webrecherche bleiben Aktualität, belastbare Quellen und korrekte Belege entscheidend; ein Wissensscore allein belegt das nicht. In Talos verbindet erst das Zusammenspiel aus Suche, Dokumentenverarbeitung, Modell und Dateierstellung diese Fähigkeiten.')
p('Coding bleibt eine ergänzende Fähigkeit: DeepSWE 1.1 nennt für Qwen 27B / Qwen Flash Next / DeepSeek Flash / GLM Flash / GLM groß die Werte 42,2 / 58,7 / 54,4 / 63,4 / 66,9. Anbieter verwenden unterschiedliche Agentenrahmen; daraus folgt keine Rangfolge für Büroarbeit. [S14–S16]')

page('6 So arbeitet die Software in Talos')
p('Der vorliegende Talos-Projektstand verbindet eine React-Oberfläche mit einem Python-Backend, getrennten Modellservern und Diensten für Dokumente und Werkzeuge. Die Tabelle zeigt die tatsächlich verwendete Software aus Compose-Datei, Abhängigkeiten, Architektur und Setup-Guide. [L1–L5]')
table(['Talos-Baustein','Verwendete Software','Konkrete Rolle'],[
('Oberfläche und Anwendung','React 19, TypeScript, Vite; FastAPI und Uvicorn','Weboberfläche; Backend steuert Chat, Dateien und Agentenschleife. Service talos-app, Port 7000.'),
('Chatmodell','vLLM; Qwen3.8-27B-NVFP4','Setup-Endpunkt 8000; Qwen3.6-35B-A3B-NVFP4 als Alternative. Talos verbindet sich über LLM_HOST(S).'),
('Suchmodelle','Qwen3-VL-Embedding-2B; Qwen3-VL-Reranker-2B','Separate Endpunkte 8001 und 8002: Suchvektoren erzeugen und Treffer neu sortieren.'),
('Dokumente und Suche','Docling, Haystack, FastEmbed; Qdrant','Dateien auslesen und zerlegen; dichte Vektorsuche plus BM25-Begriffsabgleich; Fundstellen für Antworten.'),
('Hintergrundjobs','RQ und Redis-Client; Valkey 8','Service rag-redis verwendet Valkey. rag-ingest-worker verarbeitet Uploads außerhalb der Chat-Anfrage.'),
('Daten und Gedächtnis','SQLAlchemy mit SQLite/PostgreSQL; optional Chroma','Anwendungsdaten in SQL; Chroma für Memory-/Tool-Index laut Architektur. Dokumenten-RAG verwendet Qdrant.'),
('Web und Codearbeit','SearxNG; talos-sandbox','Websuche und separate Ausführungsumgebung. Sandbox im internen Netz, Compose-Limit 6 GB RAM und 4 CPUs.'),
('Werkzeuge und Sprache','Python MCP SDK; optional Qwen3-ASR-1.7B','MCP verbindet Tools und Daten. ASR-Endpunkt 8003 übernimmt optionale Transkription.')],[1.34,2.31,3.18],9.2)
h('Beispiel einer Dokumentenfrage in Talos')
p('Upload → RQ/Valkey → Ingest-Worker → Docling und Haystack → Embeddings → Qdrant. Anschließend: Frage aus React → FastAPI → hybride Suche mit Zugriffsf ilter → Reranker → relevante Textstellen zum Qwen-Modell → Antwort mit Quellen in der Oberfläche.'.replace('Zugriffsf ilter','Zugriffsfilter'))
h('MCP am konkreten Talos Aufbau')
p('Die Agentenschleife kann als MCP-Client externe Werkzeuge aufrufen. Umgekehrt bietet Talos unter /mcp einen eigenen Server mit stateless Streamable HTTP: externe Clients können freigegebene Dokumente, Skills und Webfunktionen lesen. Dazu gehören Scopes wie rag:read, skills:read und web:read. MCP vermittelt Werkzeugaufrufe; vLLM führt das Sprachmodell aus. [L3–L4, S27]')
p('Docker Compose betreibt die Talos-Dienste; die Modellendpunkte können auf dem Spark oder einem separaten RTX-Server liegen. Ein Hardwarewechsel am Modellserver ist deshalb möglich, ohne die gesamte Talos-Anwendung umzuziehen. Der Setup-Guide beschreibt die Endpunkte; daraus folgt keine Aussage, welche Dienste gerade aktiv laufen.')

page('7 Hardware Upgrades und ihre Wirkung')
p('Die Ausbaupfade beziehen sich auf zusätzliche oder ersetzte Hardware. Entscheidend ist der verbleibende Engpass: Modellkapazität, gleichzeitige Anfragen, Antworttempo oder Datenverarbeitung.')
table(['Hardware Upgrade','Nutzen','Grenzen und Voraussetzungen'],[
('Zweiter Spark','256 GB nomineller Gesamtspeicher; größeres verteiltes Modell oder zweiter unabhängiger Modellserver.','RAM wird nicht automatisch zusammengelegt. Im Cluster entstehen Netzwerkaufwand und Abhängigkeit vom zweiten Gerät.'),
('Ausbau auf vier Sparks','512 GB nominell; mehr Modell-/Cachekapazität oder mehrere parallele Modellinstanzen.','Geeignete schnelle Verkabelung und Topologie nötig. NVIDIA Cluster Assistant sieht vier Geräte mit Switch vor. [S2]'),
('Schnelles Clusternetz','Passende ConnectX-Verbindungen, DAC-/AOC-Kabel und gegebenenfalls RoCE-fähiger Switch.','Verteilte Inferenz profitiert von geringer Latenz und hoher Bandbreite. Normales 10-GbE ist kein gleichwertiger Ersatz für das schnelle Modellnetz.'),
('RTX PRO 6000 Blackwell ergänzen','96 GB schneller VRAM je Karte; zwei/vier Karten ergeben 192/384 GB verteilt.','CPU-PCIe-Lanes, P2P, Steckplatzabstand, Gehäuse, Netzteil und Kühlung für den Endausbau auslegen.'),
('RTX 6000 Ada durch Blackwell ersetzen','48 auf 96 GB je Karte; Bandbreite von 960 auf 1.792 GB/s und Blackwell-Rechenformate. [S4, S12]','Generationswechsel der GPU; Edition und Leistungsaufnahme prüfen. Eine zusätzliche Ada-Karte bietet nicht dieselben Formate.'),
('Mehr NVMe Speicher','Platz für Modelle, Dokumente, Indizes und parallele Modellstände; schnelleres Laden bei passender SSD.','Erhöht den GPU-Speicher nicht. Austauschbarkeit und Garantie beim konkreten GB10-OEM prüfen.'),
('Separater Talos Datenhost','Zusätzliche CPU, RAM und SSD für Upload, Indexierung, Datenbank und Sandbox; entlastet den Modellrechner.','Beispielplanung: 8–16 CPU-Kerne, 32–64 GB RAM, 1–2 TB NVMe. Last und Datenbestand bestimmen den tatsächlichen Bedarf.'),
('Apple oder AMD mit mehr Speicher','Ultra mit 256/512 GB oder AMD PRO 495 mit bis 192 GB als Kapazitätsausbau.','Integrierter Speicher meist beim Kauf festgelegt; praktisch Geräteersatz. M5 Ultra am Stichtag erst angekündigt. [S6, S9]')],[1.6,2.48,2.75],9.1)
h('Passender Ausbau nach Ziel')
p('Mehr parallele Qwen-Nutzer: zusätzlichen Modellrechner oder weitere RTX-Karte vorsehen. Größere Flash-Modelle: zwei Sparks, ein ausreichend großer Ultra oder einen Mehr-GPU-Server vergleichen. Kürzere Einzelantworten: RTX PRO 6000 anhand derselben Aufgaben gegen GB10 bewerten. Volles GLM-5.3: eigene große Speicherplattform planen; vier RTX-Karten mit 384 GB bieten bei etwa 375 GB reinen 4-Bit-Gewichten zu wenig Reserve.')

page('8 Terminologie der Hardware')
table(['Begriff','Bedeutung und praktische Relevanz'],[
('CPU','Allgemeiner Prozessor. Steuert Dienste, Dateien und Datenverarbeitung; wichtig für Talos, OCR und Indexierung.'),
('GPU und NPU','GPU: paralleler Rechenbeschleuniger. NPU: spezialisierter KI-Beschleuniger. Eine hohe NPU-TOPS-Zahl sagt nicht, ob die gewünschte LLM-Laufzeit darauf läuft.'),
('SoC und GB10','System on Chip integriert mehrere Bausteine. GB10 ist der Grace-Blackwell-Chip der Spark-Klasse; DGX Spark ist das Gesamtsystem.'),
('Blackwell und Ada','NVIDIA-Architekturgenerationen. RTX 6000 Ada und RTX PRO 6000 Blackwell unterscheiden sich erheblich in Speicher und unterstützten Rechenformaten.'),
('ARM64 und x86 64','Prozessorarchitekturen. Spark und Apple verwenden Arm; viele AMD-Server und PCs x86. Container und Bibliotheken brauchen passende Versionen.'),
('RAM und VRAM','RAM ist Hauptspeicher des Systems; VRAM ist Speicher einer diskreten GPU. 512 GB CPU-RAM machen eine 24-GB-GPU nicht zu einer 512-GB-GPU.'),
('Unified Memory und UMA','CPU und GPU teilen sich physischen Speicher. Kapazität wird gemeinsam verbraucht; Betriebssystem und CPU-Aufgaben konkurrieren mit dem Modell.'),
('LPDDR GDDR und HBM','Speichertypen. LPDDR ist energieeffizient; GDDR gehört typischerweise zu Grafikkarten; HBM bietet sehr hohe Bandbreite in Serverbeschleunigern.'),
('Speicherbandbreite','Pro Sekunde transportierbare Daten, meist GB/s oder TB/s. Besonders für die fortlaufende Ausgabe dichter Modelle wichtig.'),
('TOPS FLOPS und Sparsity','TOPS zählt Billionen Operationen pro Sekunde; FLOPS Gleitkommaoperationen, PFLOPS Billiarden davon. Sparsity nutzt dünn besetzte Werte; Spitzenwerte gelten nur für passende Formate und Aufgaben.'),
('GB und GiB','1 GB = 1.000.000.000 Byte; 1 GiB = 1.073.741.824 Byte. Unterschiedliche Anzeigen sind nicht automatisch verlorener Speicher.'),
('PCIe und Lanes','Verbindung zwischen CPU und Steckkarten. Generation und Anzahl elektrischer Lanes begrenzen Datentransfer und Erweiterbarkeit.'),
('NVLink C2C und NVLink','C2C verbindet CPU und GPU innerhalb geeigneter NVIDIA-Chipsysteme. Daraus folgt keine externe NVLink-Verbindung zwischen zwei Sparks.'),
('RDMA RoCE QSFP und ConnectX','RDMA überträgt Daten mit wenig CPU-Beteiligung; RoCE nutzt Ethernet. QSFP bezeichnet Anschluss-/Modulbauformen; ConnectX die NVIDIA-Netzwerkkartenfamilie.'),
('P2P und Topologie','Peer-to-Peer ist direkter Datentransfer zwischen Beschleunigern. Topologie beschreibt ihre Verbindung; ein langsamer Umweg über die CPU kann skalierte Inferenz bremsen.'),
('ECC TDP und Netzteilleistung','ECC korrigiert bestimmte Speicherfehler. TDP ist eine thermische Auslegungsgröße. Netzteilnennleistung ist weder Dauerverbrauch noch Stromrechnung.'),
('NVMe SSD und Offloading','NVMe-SSD speichert Gewichte und Daten dauerhaft. Offloading verlagert Modellteile in RAM oder SSD und spart VRAM, kann aber die Antwort stark verlangsamen.')],[2.02,4.81],9.5)

page('9 Terminologie der Modelle und Leistung')
table(['Begriff','Bedeutung und praktische Relevanz'],[
('LLM und VLM','Large Language Model verarbeitet Sprache; Vision Language Model zusätzlich visuelle Eingaben. Bildverständnis ist keine Bildgenerierung.'),
('Parameter und B','Gelernte Zahlen des Modells. B steht für eine Milliarde. Englisches T steht für eine Billion. Mehr Parameter garantieren keine höhere Qualität.'),
('Dense und MoE','Dense aktiviert große Teile des Modells pro Token. Mixture of Experts wählt Experten aus. MoE spart aktive Arbeit, benötigt aber Speicher für die gesamten Experten.'),
('A3B und aktive Parameter','A3B bezeichnet ungefähr drei Milliarden aktive Parameter pro Token. Bei 35B-A3B müssen trotzdem rund 35B Parameter gespeichert werden.'),
('Gewichte Checkpoint und Revision','Gewichte sind die gelernten Zahlen. Checkpoint bezeichnet einen konkreten Modellstand; Revision bindet die genaue Dateiversion.'),
('Token und Kontextfenster','Token sind Text-/Datenbausteine, nicht exakt Wörter. Kontext umfasst unter anderem Prompt, Verlauf, Dokumente und Ausgabe; bei Vision entstehen weitere Tokens.'),
('Thinking und Reasoning Budget','Zusätzliche generierte Denk-Tokens beziehungsweise ihr Limit. Sie verbrauchen Zeit und Kontext, auch wenn die Oberfläche sie nicht als Antwort zeigt.'),
('BF16 FP16 FP8 INT4 und FP4','Zahlenformate mit unterschiedlicher Bitbreite. Weniger Bits sparen typischerweise Speicher; Genauigkeit und Unterstützung hängen vom Verfahren ab.'),
('Quantisierung und W4A16','Komprimierung numerischer Werte. W4A16 bedeutet 4-Bit-Gewichte und 16-Bit-Aktivierungen. Nicht alle Modellteile müssen dasselbe Format verwenden.'),
('NVFP4 GGUF Q4 und MLX 4bit','NVFP4 ist ein NVIDIA-4-Bit-Format. GGUF ist ein Dateiformat mit verschiedenen Quantisierungen wie Q4_K_M. MLX-4bit ist ein anderes Format/Laufzeitpaket.'),
('KV Cache und Zustandsbuffer','Zwischenspeicher der bisherigen Verarbeitung. Wächst je nach Architektur mit Kontext und Parallelität; lineare/hybride Modelle haben zusätzliche Zustände.'),
('Prefill Decode und TTFT','Prefill verarbeitet die Eingabe; Decode erzeugt neue Tokens. Time to First Token misst die Wartezeit bis zum ersten Token und kann auch Warteschlange enthalten.'),
('TPOT ITL und Ende zu Ende','Time per Output Token und Inter-Token Latency beschreiben Ausgabetaktung. Ende-zu-Ende-Zeit umfasst die gesamte Anfrage inklusive Warten und Werkzeugen.'),
('Durchsatz und Concurrency','Durchsatz zählt Arbeit pro Zeit; t/s bedeutet Tokens pro Sekunde. Concurrency zählt aktive parallele Anfragen. Gesamte t/s sind nicht t/s pro Nutzer.'),
('MTP und Speculative Decoding','Multi-Token Prediction oder ein Draft-Modell schlägt mehrere Tokens vor; das Zielmodell prüft sie. Häufige Ablehnung kann den Vorteil aufheben.'),
('DFlash2 DSpark und Acceptance','Spezielle Entwurfsmodelle/-verfahren. Acceptance beschreibt die akzeptierten Vorschläge, nicht die fachliche Richtigkeit der Antwort.'),
('Temperatur und Greedy','Sampling steuert die Zufälligkeit. Greedy wählt jeweils das wahrscheinlichste Token; Temperatur 0 wird oft entsprechend umgesetzt.'),
('P50 P95 und P99','Perzentile einer Messreihe. P95 von 3 Sekunden bedeutet, dass 95 % der gemessenen Fälle höchstens 3 Sekunden benötigen.'),
('QSA lineare Aufmerksamkeit und N Gramm','QSA wählt relevante Kontextblöcke aus; lineare Aufmerksamkeit verdichtet Kontext in Zuständen. N-Gram-Tabellen speichern gelernte Werte für kurze Tokenfolgen und benötigen eigenen Speicher.')],[2.04,4.79],9.4)

page('10 Terminologie der Software und Qualität')
table(['Begriff','Bedeutung und praktische Relevanz'],[
('Inference und Training','Inference wendet ein Modell an. Training verändert Gewichte. Ein Gerät, das ein quantisiertes Modell ausführt, kann es nicht automatisch vollständig trainieren.'),
('Fine Tuning LoRA und QLoRA','Anpassung durch weiteres Training; LoRA trainiert kleine Zusatzmatrizen, QLoRA arbeitet mit quantisiertem Basismodell. Dafür werden Daten und zusätzliche Speicherreserven benötigt.'),
('API Endpoint und Streaming','API ist eine Programmschnittstelle; Endpoint eine konkrete Adresse. Streaming liefert Teile der Antwort bereits während der Generierung.'),
('OpenAI kompatible API','Nachgebildetes Schnittstellenformat. Das Modell kann vollständig lokal und von einem anderen Anbieter sein; Kompatibilität ist je Funktion zu prüfen.'),
('MCP Host Client und Server','Der Host ist die KI-Anwendung, der Client verbindet sie mit einem Server. Der Server bietet Werkzeuge, Ressourcen oder Vorlagen an. [S27]'),
('Tool Calling und Parser','Das Modell erzeugt strukturierte Funktionsaufrufe; Parser übersetzen Modellausgabe in erwartete Felder. Fehler können wie schlechte Modellleistung aussehen.'),
('Skill Agent und Harness','Skill: beschriebene Vorgehensweise. Agent: Modell mit Werkzeug- und Handlungsschleife. Harness: Umgebung und Regeln, unter denen der Agent arbeitet oder bewertet wird.'),
('RAG Embedding und Reranker','Retrieval Augmented Generation ergänzt gefundene Quellen. Embedding bildet Inhalte als Suchvektoren ab; Reranker sortiert Fundstellen genauer.'),
('Chunking Hybrid Search und BM25','Chunking zerlegt Dokumente. Hybride Suche kombiniert Vektorsuche mit Begriffsabgleich, etwa BM25. Ein guter Index verbessert Antworten ohne größeres LLM.'),
('OCR und ASR','Optical Character Recognition liest Text aus Bildern. Automatic Speech Recognition wandelt Sprache in Text um; beides sind eigene Verarbeitungsschritte.'),
('vLLM SGLang und llama cpp','Programme für Modellinferenz und Serving. Unterschiedliche Architektur-, Quantisierungs- und Hardwareunterstützung; keine universelle Austauschbarkeit.'),
('CUDA ROCm Metal und MUSA','Beschleunigungsplattformen von NVIDIA, AMD, Apple und Moore Threads. CANN gehört zum Huawei-Ascend-Ökosystem.'),
('Kernel und Backend','Kernel ist eine konkrete Rechenroutine; Backend wählt deren Ausführung. Fehlende oder fehlerhafte Routinen können ein ansonsten passendes Modell unbrauchbar machen.'),
('Tensor Pipeline und Data Parallel','Tensor Parallel teilt Rechenschritte, Pipeline Parallel verteilt Modellschichten, Data Parallel betreibt Replikate für unterschiedliche Anfragen.'),
('Prefix Cache und OOM','Prefix Cache nutzt bereits verarbeitete Promptteile erneut. Out of Memory bedeutet Speichermangel. Ein warmer Cache kann Benchmarks stark beschleunigen.'),
('Container Image Digest und Rollback','Container bündeln Software. Digest bindet eine konkrete Imageversion; Rollback kehrt zu einem vorherigen Stand zurück. Persistente Daten müssen dazu kompatibel bleiben.'),
('GPQA HLE SWE und Terminal Bench','Tests für Fachwissen/Reasoning, schwere Wissensfragen, Softwareaufgaben und Terminalarbeit. Versionen und Agentenrahmen müssen bei Zahlen mitgenannt werden.'),
('Toolathlon MCPMark und Pass at 1','Werkzeugtests und ihre Erfolgsmaße. Pass@1 misst Erfolg beim ersten Versuch; ein hoher Wert garantiert keine korrekte Ausführung jedes eigenen Werkzeugs.'),
('Open Weights und Lizenz','Offene Gewichte sind herunterladbar; Nutzungsbedingungen bleiben modellabhängig. Ein Quantisierungsupload hebt die Lizenz des Ausgangsmodells nicht auf.')],[2.12,4.71],9.25)

SOURCES=[
('S1','NVIDIA DGX Spark Hardware Overview','https://docs.nvidia.com/dgx/dgx-spark/hardware.html','Herstellerspezifikationen und Leistungsaufnahme.'),
('S2','NVIDIA DGX Spark Release Notes','https://docs.nvidia.com/dgx/dgx-spark/release-notes.html','Aktueller Softwarestand sowie Juni-/Juli-Updates und Cluster Assistant.'),
('S3','ASUS Ascent GX10 FAQ','https://www.asus.com/de/support/faq/1056142/','Herstellerspezifikationen der GB10-Alternative.'),
('S4','NVIDIA RTX 6000 Ada Generation','https://www.nvidia.com/en-us/design-visualization/rtx-6000/','48-GB-Generation zur Abgrenzung von Blackwell.'),
('S5','Apple Mac Studio mit M3 Ultra','https://www.apple.com/newsroom/2025/03/apple-unveils-new-mac-studio-the-most-powerful-mac-ever/','Referenzgeneration; historische Maximalbestückung, keine aktuelle Lieferzusage.'),
('S6','Apple kündigt Mac Studio mit M5 Max und M5 Ultra an','https://www.apple.com/newsroom/2026/08/apple-introduces-new-mac-studio-with-m5-max-and-m5-ultra/','25. August 2026; angekündigte Verfügbarkeit ab 22. September.'),
('S7','AMD Ryzen AI Max plus 395','https://www.amd.com/en/products/processors/laptop/ryzen/ai-300-series/amd-ryzen-ai-max-plus-395.html','Prozessordaten; Speicherbandbreite im Text aus Datenrate und Busbreite berechnet.'),
('S8','AMD Ryzen AI Halo und Ryzen AI Max PRO 400','https://www.amd.com/en/blogs/2026/amd-powers-next-generation-agent-computers-with-new-ryzen-ai-hal.html','Ankündigung, Q3-Verfügbarkeit bei OEMs und bis 160 GB Grafikzuweisung.'),
('S9','AMD Ryzen AI Max plus PRO 495','https://www.amd.com/en/products/processors/laptop/ryzen-pro/ai-max-pro-400-series/amd-ryzen-ai-max-plus-pro-495.html','192 GB, LPDDR5x-8533 und 256-Bit-Schnittstelle.'),
('S10','Huawei Atlas 950 SuperPoD auf der WAIC 2026','https://www.huawei.com/cn/news/2026/7/atlas-950-superpod','Herstellerankündigung im Juli; chinesisch. Ergänzend GLM-Ascend-Unterstützung in S16.'),
('S11','Moore Threads MTT S5000 und aktuelle Meldungen','https://docs.mthreads.com/driver-linux-server/driver-linux-server-doc-online/MTT_S5000/introduction/','80 GB und 1,6 TB/s; Herstellerdokumentation, chinesisch.'),
('S11a','Moore Threads Nachrichtenstand','https://www.mthreads.com/?source=mdc','Hersteller-Newsindex: GLM-Flash-Anpassung 27. August und Inferenzkooperation 4. September 2026.'),
('S12','NVIDIA RTX PRO 6000 Blackwell','https://www.nvidia.com/en-us/products/workstations/professional-desktop-gpus/rtx-pro-6000/','96 GB, Bandbreite und Workstation-Leistung.'),
('S12a','NVIDIA RTX PRO 6000 Editionsvergleich','https://www.nvidia.com/en-us/products/workstations/professional-desktop-gpus/rtx-pro-6000-family/','Server, Workstation und Max-Q; Kühlung und Leistungsgrenzen.'),
('S13','AMD MI350X Plattformbroschüre','https://www.amd.com/content/dam/amd/en/documents/instinct-tech-docs/product-briefs/amd-instinct-mi350x-platform-brochure.pdf','288 GB HBM3E und 8 TB/s je Beschleuniger.'),
('S14a','Qwen3.6-35B-A3B Modellkarte','https://huggingface.co/Qwen/Qwen3.6-35B-A3B','Architektur, native Kontextlänge, Lizenz und Anbieterbenchmarks.'),
('S14b','Qwen3.8-27B Modellkarte','https://huggingface.co/Qwen/Qwen3.8-27B','Modellbeschreibung und Qualitätswerte.'),
('S14','Qwen3.8-Flash-Next Modellkarte','https://huggingface.co/Qwen/Qwen3.8-Flash-Next','125B + 51B + 4B; Architektur und Vergleichsbewertungen.'),
('S14c','Qwen Flash Next Veröffentlichung','https://qwen.ai/blog?id=qwen3.8-flash-next','26. August 2026; Architekturvorstellung.'),
('S15','GLM-5.3-Flash Modellkarte','https://huggingface.co/zai-org/GLM-5.3-Flash','320B/18B, Modalität, Lizenz und Servinghinweise.'),
('S16','GLM-5.3 Modellkarte','https://huggingface.co/zai-org/GLM-5.3','753B im Repository, Lizenz, Benchmarks und Ascend-Hinweis.'),
('S17','DeepSeek-V4-Flash Modellkarte','https://huggingface.co/deepseek-ai/DeepSeek-V4-Flash','284B/13B; Modellfamilie und Referenzgewichte. API-Updates zusätzlich S28 beachten.'),
('S18','hasso5703 GB10 Messberichte','https://github.com/hasso5703/dgx-spark-qwen38/blob/main/BENCHMARKS.md','Primärmessungen und klar gekennzeichnete Reproduktionen; Qwen 27B und Flash Next. Laufzeitabhängig.'),
('S20','elsung Spark und DeepSeek Messberichte','https://github.com/elsung/dgx-spark-deepseek-v4-flash/blob/main/results/FINAL-BENCHMARKS.md','Ein-/Zwei-Spark-Messungen; Metriken und Quantisierungsbezeichnungen kritisch getrennt.'),
('S21','GLM-Flash auf zwei Sparks','https://forums.developer.nvidia.com/t/glm-5-3-flash-running-on-2x-dgx-spark-sm-121-day-0-24-7-30-3-tok-s-with-mtp-5-two-silent-gb10-gotchas-worth-knowing/381433','Direkter Testbericht vom 27. August mit Median, Kontextnachtest und Funktionsgrenzen.'),
('S22','cfontes GLM-Flash auf vier Sparks','https://huggingface.co/cfontes/glm-5.3-flash-dflash2-tp4','Primärmessungen und verlinkte JSONs; gemischte Last getrennt von Coding-Spitze.'),
('S23','julianmb Qwen3.8 auf AMD Strix Halo','https://github.com/julianmb/q38rocm','Primärmessungen; Greedy-MTP und starke Temperaturabhängigkeit.'),
('S24','Rapid-MLX große Modelle auf M3 Ultra','https://github.com/raullenchai/Rapid-MLX/blob/main/docs/benchmarks/recent-large-models-m3-ultra.md','Gleicher 256-GB-Mac, Revisionen, Mediane und Speicherangaben; Entwicklungskandidat.'),
('S25','Vadi Taslim Qwen3.6 auf RTX PRO 6000','https://www.vaditaslim.com/blog/ai/qwen3.6-benchmarks-rtx-pro-6000','BF16-Messung; Batch 2, 30 Prompts, Einzeltestlauf.'),
('S26','Vadi Taslim Qwen3.8 auf NVIDIA und AMD','https://www.vaditaslim.com/blog/ai/qwen3.8-27b-two-rigs','17. August 2026; unterschiedliche Formate und Laufzeiten, rohe versus sichtbare Ausgabe.'),
('S27','MCP Architektur und Serverfunktionen','https://modelcontextprotocol.io/specification/2025-06-18/architecture','Protokollgrundlagen; Talos-spezifische Implementierung siehe L3.'),
('S28','DeepSeek Änderungsprotokoll','https://api-docs.deepseek.com/updates/','Flash 0731, Pro 0813 und Vision-Exp vom 21. August; Anbieterbenchmarks.'),
('S29','NVIDIA IFA 2026 und RTX Spark','https://blogs.nvidia.com/blog/local-ai-ifa-next-gen-agents-nv-pair-rtx-spark/','Aktuelle IFA-Ankündigung; RTX Spark für Oktober angekündigt.'),
('S30','AMD Radeon AI PRO R9700','https://www.amd.com/en/products/graphics/workstations/radeon-ai-pro/ai-9000-series/amd-radeon-ai-pro-r9700.html','Diskrete 32-GB-GPU als Alternative für kleinere Modelle.'),
('S31','NVIDIA NIM auf DGX Spark','https://docs.nvidia.com/nim/vision-language-models/latest/deploy-on-dgx-spark.html','Aktuelle Anleitung mit GLM-5.3-Flash; konkrete Image-/Modellkombination maßgeblich.'),
]
SOURCES += [
('AA1','Artificial Analysis Qwen3.6 35B A3B','https://artificialanalysis.ai/models/qwen3-6-35b-a3b','Reasoning; Index v4.2 und aktuelle API-Ausgaberate.'),
('AA2','Artificial Analysis Qwen3.8 27B','https://artificialanalysis.ai/models/qwen3-8-27b','xhigh; Index v4.2 und API-Messung; Methodenerklärung auf der Seite.'),
('AA3','Artificial Analysis Qwen3.8 Flash Next','https://artificialanalysis.ai/models/qwen3-8-flash-next','Index v4.2 und API-Ausgaberate.'),
('AA4','Artificial Analysis DeepSeek V4 Flash 0731','https://artificialanalysis.ai/models/deepseek-v4-flash','Reasoning max; aktueller Juli-Modellstand, Index v4.2.'),
('AA5','Artificial Analysis GLM 5.3 Flash','https://artificialanalysis.ai/models/glm-5-3-flash','Index v4.2 und API-Ausgaberate.'),
('AA6','Artificial Analysis GLM 5.3','https://artificialanalysis.ai/models/glm-5-3','Max; Index v4.2 und API-Ausgaberate. Alle AA-Werte sind dynamische Momentaufnahmen.')]
SOURCES += [
('AA7','Artificial Analysis GDPval AA v2','https://artificialanalysis.ai/evaluations/gdpval-aa','Aktuelles Elo-Leaderboard und Berufsaufgaben; Werte können sich bei neuen Vergleichen ändern.'),
('AA8','Artificial Analysis AA Briefcase','https://artificialanalysis.ai/evaluations/aa-briefcase','Wissensarbeit, Memos, Tabellen und Präsentationen; aktuelle Elo-Werte und Unsicherheitsintervalle.'),
('AA9','Artificial Analysis Long Context Reasoning','https://artificialanalysis.ai/evaluations/artificial-analysis-long-context-reasoning','Informationen aus langen Dokumenten extrahieren, verknüpfen und zusammenführen.'),
('AA10','Artificial Analysis GDP pdf','https://artificialanalysis.ai/evaluations/gdp-pdf','Detailliertes Verständnis professioneller Dokumente; keine reine Layoutbewertung.'),
('AA11','Artificial Analysis Omniscience','https://artificialanalysis.ai/evaluations/omniscience','Faktenwissen und Halluzinationen in unterschiedlichen Wissensgebieten.'),
('AA12','Artificial Analysis AnalystAgent','https://artificialanalysis.ai/evaluations/aa-analyst-agent','Quantitative Arbeit mit Dokumenten und Tabellen; pass^5 misst Erfolg in allen fünf Versuchen.')]
# Remove sources no longer cited after shortening the chapters.
SOURCES = [x for x in SOURCES if x[0] not in {'S14a','S14b','S14c','S17','S25'}]
for start in range(0,len(SOURCES),12):
 page('11 Quellen' if start==0 else 'Quellen Fortsetzung')
 if start==0:p('Alle Webquellen wurden zum Stichtag 7. September 2026 recherchiert. Herstellerdaten und Modellkarten belegen Spezifikationen beziehungsweise Anbieterbewertungen. Die benannten Entwicklerberichte belegen Fremdmessungen; sie sind keine unabhängige gemeinsame Laborprüfung. Dynamische Repositories können sich später ändern.')
 for key,title,url,note in SOURCES[start:start+12]:
  par=p();par.paragraph_format.space_after=Pt(2);r=par.add_run('['+key+'] ');r.bold=True;link(par,title,url)
  par=p(note);par.paragraph_format.space_after=Pt(9);par.runs[0].font.size=Pt(9)
 if start+12>=len(SOURCES):
  h('Lokale Projektgrundlagen')
  for t in ['[L1] output/Talos_Setup_Guide.md: vorhandene Modellendpunkte und Einrichtung.','[L2] docker-compose.yml und requirements.txt: tatsächlich deklarierte Dienste und Softwareabhängigkeiten.','[L3] docs/backend/mcp-server.md: MCP-Richtung, Transport, Berechtigungen und offene LAN-Vorgaben.','[L4] docs/architecture/overview.md: FastAPI, React, Stores und Agentenablauf.','[L5] docs/architecture/rag-pipeline.md und web/package.json: konkrete Suchpipeline und Frontend-Abhängigkeiten.']:
   par=p(t);par.runs[0].font.size=Pt(9)
doc.core_properties.title='Talos Hardware und Modelle'
doc.core_properties.subject='Hardware Software Modellvergleich und Hardware Upgrades Stand September 2026'
doc.core_properties.author=''
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
