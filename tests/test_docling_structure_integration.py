"""Actual Docling-Haystack conversion without OCR/model downloads."""

import pytest

pytest.importorskip("docling")

from src.rag_vector import VectorRAG


def rag(monkeypatch):
    instance = object.__new__(VectorRAG)
    instance._docling = None
    instance._cfg = {"chunk_max_chars": 1000, "chunk_overlap_chars": 100}
    monkeypatch.setenv("PDF_VLM_ENABLED", "")
    monkeypatch.setenv("RAG_REDACT_PII", "")
    monkeypatch.setattr(instance, "_apply_contextual", lambda _: None)
    monkeypatch.setattr(instance, "_apply_autokeywords", lambda _: None)
    return instance


@pytest.mark.parametrize("extension", ["html", "docx"])
def test_actual_converter_keeps_chapters_and_table_headers(tmp_path, monkeypatch, extension):
    text = "Hydraulic pump adjustment and pressure calibration. " * 80
    path = tmp_path / f"manual.{extension}"
    if extension == "html":
        rows = "".join(f"<tr><td>valve{i}</td><td>pressure{i}</td></tr>" for i in range(100))
        path.write_text(
            f"<html><body><h1>Manual</h1><p>Introduction.</p><h2>Setup</h2><p>{text}</p>"
            f"<h2>Values</h2><table><tr><th>Name</th><th>Pressure</th></tr>{rows}</table></body></html>",
            encoding="utf-8",
        )
    else:
        from docx import Document

        native = Document()
        native.add_heading("Manual", level=1)
        native.add_paragraph("Introduction.")
        native.add_heading("Setup", level=2)
        native.add_paragraph(text)
        native.add_heading("Values", level=2)
        table = native.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Pressure"
        for i in range(100):
            row = table.add_row()
            row.cells[0].text, row.cells[1].text = f"valve{i}", f"pressure{i}"
        native.save(path)
    docs = rag(monkeypatch)._documents_for_file(str(path), {"source": str(path)})
    assert all(len(d.content) <= 1000 for d in docs)
    setup = [d for d in docs if d.meta["heading_path"][-1:] == ["Setup"]]
    assert len(setup) > 1
    assert len({d.meta["section_id"] for d in setup}) == 1
    tables = [d for d in docs if d.meta.get("table_header")]
    assert len(tables) > 1
    assert all("Name" in d.content and "Pressure" in d.content for d in tables)
    assert "valve99" in "".join(d.content for d in tables)


def test_actual_csv_converter_retains_last_row(tmp_path, monkeypatch):
    path = tmp_path / "values.csv"
    path.write_text(
        "Name,Pressure\n" + "".join(f"valve{i},pressure{i}\n" for i in range(150)), encoding="utf-8"
    )
    docs = rag(monkeypatch)._documents_for_file(str(path), {"source": str(path)})
    assert len(docs) > 1
    assert all(len(d.content) <= 1000 for d in docs)
    assert all("Name" in d.content and "Pressure" in d.content for d in docs)
    assert "valve149" in "".join(d.content for d in docs)
